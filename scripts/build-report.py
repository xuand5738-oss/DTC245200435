"""Build editable practice report from the maintained Markdown, no screenshots."""
from pathlib import Path
import re
from docx import Document
from docx.shared import Cm, Pt, RGBColor
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

root=Path(__file__).resolve().parent.parent
text=(root/'docs/bao-cao-thuc-hanh.md').read_text(encoding='utf-8')
doc=Document()
section=doc.sections[0]
section.page_width=Cm(21); section.page_height=Cm(29.7)
section.top_margin=Cm(2); section.bottom_margin=Cm(2)
section.left_margin=Cm(3); section.right_margin=Cm(2)
normal=doc.styles['Normal']; normal.font.name='Times New Roman'; normal.font.size=Pt(13)
normal.paragraph_format.line_spacing=1.3
normal.paragraph_format.space_after=Pt(7)
for name in ['Title','Heading 1','Heading 2','Heading 3']:
    style=doc.styles[name];style.font.name='Times New Roman';style.font.color.rgb=RGBColor(0,0,0)
    style.font.size=Pt(18 if name=='Title' else 15 if name=='Heading 1' else 13)

def clean(s): return s.replace('**','').replace('`','').replace('→',' -> ')

title=doc.add_paragraph('BÁO CÁO THỰC HÀNH',style='Title');title.alignment=1
for value in ['Triển khai và Quản trị Hệ thống Phần mềm','Đề 11 Hệ thống đặt vé xem phim và sự kiện','Họ và tên: [Sinh viên điền]','Mã số sinh viên: [Sinh viên xác nhận]','Lớp: [Sinh viên điền]','Giảng viên: [Sinh viên điền]','Ngày thực hành: 24/10/2026']:
    p=doc.add_paragraph(value);p.alignment=1
doc.add_paragraph('Báo cáo trình bày cách dựng và kiểm tra ứng dụng đặt vé, cơ sở dữ liệu, reverse proxy, giám sát, log và bảo mật trong môi trường WSL. Sinh viên bổ sung ảnh minh chứng do mình tự chụp vào các mục được hướng dẫn.')
chapters=re.split(r'^## ',text,flags=re.M)[1:]
extras={
1:['Mục tiêu học tập là hiểu mỗi dịch vụ làm nhiệm vụ gì, đọc được cấu hình Compose và tự xác minh hệ thống đang hoạt động. Mỗi kết quả cần đối chiếu với cấu hình và ảnh minh chứng; thanh toán giả lập chỉ thay đổi trạng thái vé, không có giao dịch tiền thật.'],
2:['Quy trình: mở PowerShell tại D:\\DTC245200435, chạy start-wsl.ps1 và giữ cửa sổ Ubuntu mở. Cờ -d của Docker Compose chỉ đưa container chạy nền; trên máy thực hành này vẫn cần một phiên WSL hoạt động. Không chạy down -v vì tùy chọn này xóa volume dữ liệu.'],
3:['Luồng giám sát: Prometheus -> exporter -> dịch vụ cần đo. Grafana truy vấn Prometheus để vẽ biểu đồ. Luồng log: Nginx -> file trong volume -> Promtail -> Loki -> Grafana Explore. Hai luồng bổ sung nhau: biểu đồ chỉ ra sự thay đổi, còn log cho biết yêu cầu cụ thể.'],
4:['Các mốc Git: 201ea21 cấu hình Nginx; ca5ed5a giám sát; 3d26c86 log tập trung. Commit 379794e lưu script demo tái tạo được. File .env chứa bí mật chỉ nằm tại máy, còn .env.example mô tả tên biến để dựng lại. Không chụp nội dung .env vào báo cáo.'],
6:['Để kiểm tra, chạy docker compose exec nginx nginx -t. Sau đó chạy curl -I http://localhost:8280/. Website trả 200 qua Nginx và có các header cấu hình. Các header giúp trình duyệt xử lý nội dung an toàn hơn nhưng không thay thế xác thực người dùng hay mã hóa HTTPS.'],
7:['Khi dashboard chưa hiện tốc độ, đợi ít nhất hai lần scrape để rate có đủ điểm. Chọn Last 15 minutes và truy cập website để sinh request. Một target UP chỉ xác nhận endpoint scrape hoạt động; cần xem nginx_up/mysql_up và các metrics cụ thể để xác nhận dịch vụ phía sau.'],
8:['Log chứa thời gian, phương thức, URI, mã HTTP và thời gian xử lý. Nhãn project phân biệt hệ thống này với hệ thống khác. Query status=404 giúp tìm đường dẫn lỗi; query status=200 xác nhận các yêu cầu thành công. Khi sinh log mới cần đợi Promtail chuyển dữ liệu trước khi refresh.'],
9:['Lệnh kiểm tra: docker compose exec -T app node < scripts/verify-hardening.js. Kết quả phải có non-root uid=1000 và CREATE TABLE denied. Docker inspect kiểm tra ReadonlyRootfs, CapDrop và SecurityOpt. Việc bảo mật cần dựa trên quyền thực tế, không chỉ dựa vào tên tài khoản.'],
10:['Kịch bản demo: mở suất chiếu -> mở vé A1 -> đối chiếu bookings trong phpMyAdmin -> mở Prometheus targets -> mở dashboard -> chạy ba LogQL -> giải thích hardening. Không thêm vé mới nếu chỉ cần chứng minh bản ghi demo đã kiểm tra.'],
11:['Giới hạn và việc sinh viên cần hoàn thiện: điền bìa, tự chụp và chèn ảnh, kiểm tra bố cục trong Word, chuẩn bị trả lời vấn đáp. Tài khoản GitHub hiện là xuand5738-oss; yêu cầu đặt tên theo MSSV cần sinh viên xác nhận hoặc xử lý với giảng viên. Không tự đổi tài khoản trong quá trình triển khai.']}
for idx,chapter in enumerate(chapters,1):
    doc.add_page_break()
    lines=chapter.strip().splitlines();doc.add_heading(clean(lines[0]),1)
    in_code=False
    for line in lines[1:]:
        if line.startswith('```'): in_code=not in_code; continue
        if not line.strip():continue
        if line.startswith('### '):doc.add_heading(clean(line[4:]),2)
        else:
            p=doc.add_paragraph(clean(line))
            if in_code:
                for run in p.runs:run.font.name='Consolas';run.font.size=Pt(10)
    for para in extras.get(idx,[]):doc.add_paragraph(para)
    if idx in [2,4,5,6,7,8,9,10]:
        doc.add_heading('Minh chứng sinh viên bổ sung',2)
        doc.add_paragraph('[Chèn ảnh do sinh viên tự chụp và ghi chú thích nêu rõ lệnh hoặc màn hình, kết quả quan sát và tiêu chí được chứng minh.]')
footer=section.footer.paragraphs[0];footer.alignment=1
field=OxmlElement('w:fldSimple');field.set(qn('w:instr'),'PAGE');footer._p.append(field)
out=root/'docs/bao-cao-de-11.docx';doc.save(out)
assert len(chapters)==11
print('DOCX created: cover plus 11 chapters, 11 explicit page breaks; actual pagination requires Word preview.')
